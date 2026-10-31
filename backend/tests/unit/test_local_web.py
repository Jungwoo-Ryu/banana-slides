"""Local web contracts: credentials, TLS, adapters and actual PPTX artifacts."""
import base64
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock
import pytest
from PIL import Image, ImageDraw, ImageFont


def test_settings_do_not_store_or_return_credentials(client, monkeypatch):
    from models import Settings
    monkeypatch.setenv('AZURE_OPENAI_API_KEY', 'private-test-value')
    monkeypatch.setenv('AZURE_OPENAI_ENDPOINT', 'https://example.openai.azure.com')
    response = client.get('/api/settings')
    assert response.status_code == 200
    assert 'private-test-value' not in response.text
    assert 'api_key' not in Settings.__table__.columns
    assert client.put('/api/settings', json={'api_key': 'must-not-save'}).status_code == 400
    assert client.put('/api/settings', json={'output_language': 'ko'}).status_code == 200
    assert client.get('/api/settings').json['data']['output_language'] == 'ko'


def test_blank_project_and_saved_preferences_need_no_key(client, monkeypatch):
    for name in ('AZURE_OPENAI_API_KEY', 'TEXT_API_KEY', 'AZURE_OPENAI_API_KEY_FILE', 'TEXT_API_KEY_FILE'):
        monkeypatch.delenv(name, raising=False)
    assert client.get('/health').status_code == 200
    result = client.post('/api/projects', json={'idea_prompt':'한국어 발표', 'creation_type':'idea'})
    assert result.status_code in (200,201)
    assert client.post('/api/settings/verify').json['data']['available'] is False
    assert client.post('/api/openai/oauth/start').status_code == 404


def test_secret_file_and_verified_transport(tmp_path, monkeypatch):
    from services.runtime_config import secret
    from services.http_transport import http_client, validate_endpoint
    keyfile=tmp_path/'key'; keyfile.write_text('secret-file-value\n')
    monkeypatch.setenv('TEXT_API_KEY_FILE', str(keyfile))
    assert secret('TEXT_API_KEY') == 'secret-file-value'
    import ssl
    with http_client() as client:
        assert client._transport._pool._ssl_context.verify_mode == ssl.CERT_REQUIRED
    for invalid in ('http://external.example', 'https://user:pass@example.com'):
        with pytest.raises(ValueError): validate_endpoint(invalid)
    assert validate_endpoint('http://127.0.0.1:9000/') == 'http://127.0.0.1:9000'


def test_azure_deployment_used_for_text_vision_and_image_edits(tmp_path, monkeypatch):
    from services.ai_providers.text.azure_provider import AzureTextProvider
    from services.ai_providers.image.azure_provider import AzureImageProvider
    text = AzureTextProvider('unit-test-key','https://example.openai.azure.com','deployment-alias','2025-04-01-preview')
    text.client.close(); text.client=MagicMock()
    text.client.chat.completions.create.return_value=SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content='확인'))])
    path=tmp_path/'slide.png'; Image.new('RGB',(512,512),'white').save(path)
    assert text.generate_text('hello')=='확인'
    text.generate_with_image('describe',str(path))
    for call in text.client.chat.completions.create.call_args_list:
        assert call.kwargs['model']=='deployment-alias'
        assert 'max_completion_tokens' in call.kwargs
    image=AzureImageProvider('unit-test-key','https://example.openai.azure.com','image-alias','2025-04-01-preview')
    image.client.close(); image.client=MagicMock()
    payload=SimpleNamespace(data=[SimpleNamespace(b64_json=base64.b64encode(path.read_bytes()).decode())])
    image.client.images.generate.return_value=payload
    image.client.images.edit.return_value=payload
    assert image.generate_image('generate',resolution='1K').size==(512,512)
    assert image.generate_image('change selected region',ref_images=[Image.open(path)],resolution='1K').size==(512,512)
    assert image.client.images.generate.call_args.kwargs['model']=='image-alias'
    assert image.client.images.edit.call_args.kwargs['model']=='image-alias'
    assert image.client.images.edit.call_args.kwargs['image']


def test_azure_image_requires_separate_deployment(app, monkeypatch):
    from services.ai_providers import _get_model_type_provider_config
    monkeypatch.setenv('AZURE_OPENAI_API_KEY','unit-key')
    monkeypatch.delenv('AZURE_OPENAI_IMAGE_DEPLOYMENT',raising=False)
    with app.app_context():
        app.config.update(AI_PROVIDER_FORMAT='azure',IMAGE_MODEL_SOURCE='',AZURE_OPENAI_ENDPOINT='https://example.openai.azure.com',AZURE_OPENAI_IMAGE_DEPLOYMENT='')
        with pytest.raises(ValueError,match='AZURE_OPENAI_IMAGE_DEPLOYMENT'):
            _get_model_type_provider_config('image')


def test_local_parser_and_path_boundary(tmp_path, monkeypatch):
    from services.document_providers import LocalDocumentProvider
    from services.storage import LocalFileStorage
    from pptx import Presentation
    document=tmp_path/'sample.pptx'; presentation=Presentation(); slide=presentation.slides.add_slide(presentation.slide_layouts[1]); slide.shapes.title.text='한국어 제목'; presentation.save(document)
    assert '한국어 제목' in LocalDocumentProvider().parse(str(document),document.name)
    storage=LocalFileStorage(tmp_path/'uploads')
    with pytest.raises(ValueError): storage.resolve('../uploads-other/private.txt')
    assert storage.resolve('project/image.png').is_relative_to(storage.root)


def test_local_ocr_editable_pptx_roundtrip(app, tmp_path, monkeypatch):
    from services.document_providers import LocalOCRProvider
    from services.image_editability.local_adapters import LocalTextAttributeExtractor
    from services.export_service import ExportService
    from pptx import Presentation
    import pytesseract
    # The CI fixture uses English so it is independent of an optional Korean pack.
    assert 'eng' in pytesseract.get_languages()
    image=Image.new('RGB',(960,540),'white')
    font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',52) if Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf').exists() else ImageFont.truetype('arial.ttf',52)
    ImageDraw.Draw(image).text((80,150),'LOCAL EDITABLE SLIDE',fill='black',font=font)
    path=tmp_path/'slide.png'; image.save(path)
    with app.app_context():
        app.config.update(OCR_PROVIDER='local',OCR_LANGUAGES='eng',INPAINTING_PROVIDER='local',UPLOAD_FOLDER=str(tmp_path))
        # Extraction runs on worker threads outside Flask context.
        monkeypatch.setenv('OCR_LANGUAGES','eng'); monkeypatch.setenv('UPLOAD_FOLDER',str(tmp_path))
        elements=LocalOCRProvider().recognize(image)
        assert any('EDITABLE' in item['content'] for item in elements)
        target=tmp_path/'editable.pptx'
        ExportService.create_editable_pptx_with_recursive_analysis(image_paths=[str(path)], output_file=str(target), max_workers=1, text_attribute_extractor=LocalTextAttributeExtractor(), export_inpaint_method='local', enable_icon_subject_extraction=False, fail_fast=True)
    deck=Presentation(target)
    assert len(deck.slides)==1
    assert any('EDITABLE' in shape.text for shape in deck.slides[0].shapes if shape.has_text_frame)


def test_standard_export_is_real_pptx(client, tmp_path):
    from models import db, Project, Page
    from services.file_service import FileService
    from flask import current_app
    from pptx import Presentation
    project=Project(idea_prompt='local export'); db.session.add(project);db.session.flush()
    page=Page(project_id=project.id,order_index=0,status='COMPLETED');db.session.add(page);db.session.flush()
    page.generated_image_path=FileService(current_app.config['UPLOAD_FOLDER']).save_generated_image(Image.new('RGB',(960,540),'navy'),project.id,page.id)
    db.session.commit()
    result=client.get(f'/api/projects/{project.id}/export/pptx')
    assert result.status_code==200, result.json
    downloaded=client.get(result.json['data']['download_url'])
    deck=Presentation(BytesIO(downloaded.data))
    assert len(deck.slides)==1
    assert len(deck.slides[0].shapes)==1
